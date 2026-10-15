import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import import_catalog


class CatalogAtomicTests(unittest.TestCase):
    def test_exact_sku_corrections_survive_source_workbook_import(self):
        with tempfile.TemporaryDirectory() as temp:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            rows = [
                ('28308779', '28308779 A651070328 FUTURE-CODE', 'Топливная форсунка Common Rail Delphi', 'Delphi', 'Mercedes-Benz'),
                ('28540276', '28540276 R00203D', 'Топливная форсунка Common Rail Delphi', 'Delphi', 'Mercedes-Benz'),
                ('DZ111136', 'DZ111136', 'Клапан управления подачей топлива ТНВД John Deere', 'John Deere', ''),
            ]
            rows.extend((f'TEST-{i:02}', f'TEST-{i:02}', 'Деталь для проверки импорта', 'Other', '') for i in range(55))
            for row_index, (sku, codes, name, brand, cross_brand) in enumerate(rows, 7):
                sheet.cell(row_index, 3, codes)
                sheet.cell(row_index, 5, 100)
                sheet.cell(row_index, 9, name)
                sheet.cell(row_index, 11, brand)
                sheet.cell(row_index, 12, cross_brand)
                sheet.cell(row_index, 13, f'https://example.test/{sku}')
            workbook_path = Path(temp) / 'catalog.xlsx'
            workbook.save(workbook_path)

            products = {item['sku']: item for item in import_catalog.parse_catalog(workbook_path)}

            self.assertEqual(products['28308779']['codes'], ['28308779', 'A6510703287', 'FUTURE-CODE'])
            self.assertEqual(products['28540276']['crossBrand'], 'Hyundai / Kia')
            self.assertEqual(products['DZ111136']['category'], 'valves')

    def test_failed_products_publish_preserves_previous_catalog(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'data').mkdir()
            target = root / 'data/products.json'
            previous = '[{"id":"last-good"}]\n'
            target.write_text(previous, encoding='utf-8')
            products = [{'id': 'last-good'}]

            original_write_text = Path.write_text
            def fail_direct_target_write(path, data, *args, **kwargs):
                if path == target:
                    original_write_text(path, '[partial', *args, **kwargs)
                    raise OSError('injected interrupted catalog write')
                return original_write_text(path, data, *args, **kwargs)

            original_replace = os.replace
            def fail_atomic_publish(source, destination):
                if Path(destination) == target:
                    raise OSError('injected atomic catalog publish failure')
                return original_replace(source, destination)

            with patch.object(import_catalog, 'ROOT', root), \
                 patch.object(import_catalog, 'parse_catalog', return_value=products), \
                 patch.object(sys, 'argv', ['import_catalog.py', 'catalog.xlsx']), \
                 patch.object(Path, 'write_text', fail_direct_target_write), \
                 patch.object(os, 'replace', side_effect=fail_atomic_publish):
                with self.assertRaises(OSError):
                    import_catalog.main()

            self.assertEqual(target.read_text(encoding='utf-8'), previous)

    def test_import_rejects_missing_ids_without_replacing_expanded_catalog(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'data').mkdir()
            target = root / 'data/products.json'
            previous = '[{"id":"last-good"},{"id":"new-john-deere"}]\n'
            target.write_text(previous, encoding='utf-8')
            products = [{'id': 'last-good'}]

            with patch.object(import_catalog, 'ROOT', root), \
                 patch.object(import_catalog, 'parse_catalog', return_value=products), \
                 patch.object(sys, 'argv', ['import_catalog.py', 'legacy-catalog.xlsx']):
                with self.assertRaisesRegex(ValueError, 'new-john-deere'):
                    import_catalog.main()

            self.assertEqual(target.read_text(encoding='utf-8'), previous)


if __name__ == '__main__':
    unittest.main()
